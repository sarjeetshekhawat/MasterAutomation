"""Inspect template workbooks for sheets, headers, and field types."""

from __future__ import annotations

import logging
from dataclasses import dataclass, field
from typing import Dict, List, Set

import openpyxl

from .config import (
    MANDATORY_HIERARCHY_FIELDS,
    MANDATORY_TARGET_FIELDS,
    REFERENCE_SHEET_NAME,
    TARGET_FIELD_TYPES,
    UUID_FIELDS,
    FieldType,
)
from .file_loader import get_sheet_headers, get_type_hints

logger = logging.getLogger(__name__)

HIERARCHY_HEADER_NAMES = {
    "discom_name", "zone_name", "zone_code",
    "circle_name", "circle_code",
    "division_name", "division_code",
    "subdivision_name", "subdivision_code",
    "substation_name", "substation_code",
    "feeder_name", "feeder_code",
    "dt_name", "dt_code",
}


@dataclass
class FieldSchema:
    name: str
    position: int
    type_hint: str = ""
    is_uuid: bool = False
    is_mandatory: bool = False
    is_hierarchy: bool = False
    is_ref_lookup: bool = False
    field_type: FieldType = FieldType.TEXT


@dataclass
class TemplateSchema:
    sheet_name: str
    headers: List[str]
    fields: Dict[str, FieldSchema]
    has_type_hint_row: bool = False
    reference_sheet_exists: bool = False
    reference_categories: Set[str] = field(default_factory=set)

    @property
    def header_count(self) -> int:
        return len(self.headers)

    @property
    def uuid_fields(self) -> List[str]:
        return [f.name for f in self.fields.values() if f.is_uuid]

    @property
    def mandatory_fields(self) -> List[str]:
        return [f.name for f in self.fields.values() if f.is_mandatory]

    @property
    def hierarchy_fields(self) -> List[str]:
        return [f.name for f in self.fields.values() if f.is_hierarchy]

    @property
    def ref_lookup_fields(self) -> List[str]:
        return [f.name for f in self.fields.values() if f.is_ref_lookup]


def detect_template_schema(wb: openpyxl.Workbook, target_sheet: str) -> TemplateSchema:
    headers = get_sheet_headers(wb, target_sheet)
    type_hints = get_type_hints(wb, target_sheet)
    ref_exists = REFERENCE_SHEET_NAME in wb.sheetnames

    has_type_hints = False
    hint_keywords = ["uuid", "text", "double", "mandatory", "ref_"]
    sample_hints = list(type_hints.values())[:10]
    for hint in sample_hints:
        if any(kw in hint.lower() for kw in hint_keywords):
            has_type_hints = True
            break

    fields: Dict[str, FieldSchema] = {}
    for idx, header in enumerate(headers):
        if not header:
            continue
        header_lower = header.lower().strip()
        hint = type_hints.get(header, "")
        hint_lower = hint.lower()
        is_uuid = (
            header_lower in UUID_FIELDS
            or "uuid" in hint_lower
            or "ref_level_entries" in hint_lower
            or "ref_project_masters_lov" in hint_lower
        )
        is_mandatory = (
            header_lower in MANDATORY_TARGET_FIELDS
            or "mandatory" in hint_lower
        )
        is_hierarchy = header_lower in HIERARCHY_HEADER_NAMES
        is_ref = "ref_project_masters_lov" in hint_lower
        ft = TARGET_FIELD_TYPES.get(header_lower, FieldType.TEXT)
        fields[header_lower] = FieldSchema(
            name=header,
            position=idx + 1,
            type_hint=hint,
            is_uuid=is_uuid,
            is_mandatory=is_mandatory or header_lower in MANDATORY_HIERARCHY_FIELDS,
            is_hierarchy=is_hierarchy,
            is_ref_lookup=is_ref,
            field_type=ft,
        )

    ref_categories: Set[str] = set()
    if ref_exists:
        ref_ws = wb[REFERENCE_SHEET_NAME]
        headers_ref = [c.value for c in next(ref_ws.iter_rows(max_row=1))]
        master_col = 4
        for i, h in enumerate(headers_ref, 1):
            if h and str(h).strip().lower() == "master_name":
                master_col = i
                break
        for row in range(2, ref_ws.max_row + 1):
            master_name = ref_ws.cell(row=row, column=master_col).value
            if master_name:
                ref_categories.add(str(master_name).strip())

    schema = TemplateSchema(
        sheet_name=target_sheet,
        headers=headers,
        fields=fields,
        has_type_hint_row=has_type_hints,
        reference_sheet_exists=ref_exists,
        reference_categories=ref_categories,
    )
    logger.info(
        "Detected template schema: %s columns, %s UUID fields, %s mandatory, ref sheet: %s",
        schema.header_count,
        len(schema.uuid_fields),
        len(schema.mandatory_fields),
        ref_exists,
    )
    return schema


def detect_source_schema(headers: List[str]) -> Dict[str, int]:
    schema = {}
    for idx, header in enumerate(headers):
        if header:
            schema[header.strip().lower()] = idx
    return schema
